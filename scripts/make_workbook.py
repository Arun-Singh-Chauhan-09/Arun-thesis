#!/usr/bin/env python3
"""
Builds the pilot results workbook from results/*.json.

Layout follows the sheet already set up by hand: rows KPI1..KPI6 down column A,
with 'output', 'if Issue' and 'Summary' across columns B, C and D. One sheet per
sample, plus a Summary sheet and a Findings sheet.

Usage:
    python3 scripts/make_workbook.py                       # from results/
    python3 scripts/make_workbook.py --template            # empty, for testing
"""
import argparse
import json
import pathlib

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = pathlib.Path(__file__).resolve().parent.parent
KPIS = [
    ("KPI1", "Is it parseable YAML?", "PyYAML"),
    ("KPI2", "Is it a valid Kubernetes object?", "kubeconform v1.31"),
    ("KPI3", "Does the cluster accept it?", "kubectl apply (kind)"),
    ("KPI4", "Does it become ready?", "kubectl wait, 180s"),
    ("KPI5", "Does it do what was asked, and no more?", "intent probe"),
    ("KPI6", "Free of HIGH/CRITICAL faults?", "Checkov + Trivy + kube-linter"),
]

HDR_FILL = PatternFill("solid", fgColor="D9D9D9")
PASS_FILL = PatternFill("solid", fgColor="E2EFDA")
FAIL_FILL = PatternFill("solid", fgColor="FCE4E4")
NA_FILL = PatternFill("solid", fgColor="FFF2CC")
INPUT_FILL = PatternFill("solid", fgColor="FFFFCC")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BASE = Font(name="Arial", size=10)
BOLD = Font(name="Arial", size=10, bold=True)
TITLE = Font(name="Arial", size=12, bold=True)


def style_cell(c, font=BASE, fill=None, wrap=False, align="left"):
    c.font = font
    c.border = BORDER
    c.alignment = Alignment(horizontal=align, vertical="top", wrap_text=wrap)
    if fill:
        c.fill = fill


def v_to_text(v):
    if v is True:
        return "PASS"
    if v is False:
        return "FAIL"
    if v == "N/A":
        return "N/A"
    return ""


def fill_for(text):
    return {"PASS": PASS_FILL, "FAIL": FAIL_FILL, "N/A": NA_FILL}.get(text, INPUT_FILL)


def run_sheet(wb, title, result):
    ws = wb.create_sheet(title)
    ws.column_dimensions["A"].width = 9
    ws.column_dimensions["B"].width = 11
    ws.column_dimensions["C"].width = 62
    ws.column_dimensions["D"].width = 40
    ws.column_dimensions["E"].width = 30

    for col, head in zip("BCDE", ["output", "if Issue", "Summary", "Decided by"]):
        c = ws[f"{col}1"]
        c.value = head
        style_cell(c, BOLD, HDR_FILL)
    style_cell(ws["A1"], BOLD, HDR_FILL)

    verdict = (result or {}).get("verdict", {})
    detail = (result or {}).get("detail", {})

    for i, (kpi, question, tool) in enumerate(KPIS, start=2):
        key = kpi.replace("KPI", "K")
        out = v_to_text(verdict.get(key))
        issue = detail.get(key, "") if out in ("FAIL", "N/A") else ""
        summary = detail.get(key, "") if out == "PASS" else question

        style_cell(ws.cell(i, 1, kpi), BOLD)
        style_cell(ws.cell(i, 2, out), BOLD, fill_for(out), align="center")
        style_cell(ws.cell(i, 3, issue), wrap=True)
        style_cell(ws.cell(i, 4, summary), wrap=True)
        style_cell(ws.cell(i, 5, tool))
        ws.row_dimensions[i].height = 30

    style_cell(ws.cell(9, 1, "Class"), BOLD)
    ws["B9"] = ('=IF(COUNTA(B2:B7)<6,"",'
                'IF(OR(B2="FAIL",B3="FAIL"),"INVALID",'
                'IF(COUNTIF(B4:B7,"FAIL")>0,"MISCONFIGURED","CLEAN")))')
    style_cell(ws["B9"], BOLD, HDR_FILL, align="center")
    style_cell(ws.cell(9, 3, "INVALID = fails K1/K2 - MISCONFIGURED = fails any of K3-K6 - "
                             "CLEAN = passes all six"), wrap=True)

    meta = [
        ("Generation ID", (result or {}).get("generation_id", "")),
        ("Scenario", (result or {}).get("scenario", "S01")),
        ("Condition", "P0 (zero-shot)"),
        ("Model", "M1 - deepseek-chat"),
        ("Namespace", (result or {}).get("namespace", "")),
        ("Evaluated (UTC)", (result or {}).get("evaluated_at_utc", "")),
        ("Notes", "; ".join((result or {}).get("notes", []))),
    ]
    r = 11
    style_cell(ws.cell(r, 1, "Run metadata"), BOLD, HDR_FILL)
    for label, value in meta:
        r += 1
        style_cell(ws.cell(r, 1, label), BOLD)
        style_cell(ws.cell(r, 2, value), wrap=True)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
    return ws


def summary_sheet(wb, sheet_names):
    ws = wb.create_sheet("Summary", 0)
    ws.column_dimensions["A"].width = 26
    for col in "BCD":
        ws.column_dimensions[col].width = 18
    ws.column_dimensions["E"].width = 46

    style_cell(ws.cell(1, 1, "S01 / P0 / M1 - pilot results"), TITLE)
    ws.merge_cells("A1:E1")

    style_cell(ws.cell(3, 1, "Sample"), BOLD, HDR_FILL)
    style_cell(ws.cell(3, 2, "Classification"), BOLD, HDR_FILL)
    style_cell(ws.cell(3, 3, "Failed at"), BOLD, HDR_FILL)

    for i, name in enumerate(sheet_names, start=4):
        style_cell(ws.cell(i, 1, name), BOLD)
        ws.cell(i, 2, f"='{name}'!B9")
        style_cell(ws.cell(i, 2), BASE, align="center")
        ws.cell(i, 3, f'=IF(COUNTIF(\'{name}\'!B2:B7,"FAIL")=0,"-",'
                      f'INDEX(\'{name}\'!A2:A7,MATCH("FAIL",\'{name}\'!B2:B7,0)))')
        style_cell(ws.cell(i, 3), BASE, align="center")

    last = 3 + len(sheet_names)
    rows = []
    rows.append(("CLEAN count", f'=COUNTIF(B4:B{last},"CLEAN")',
                 "Number of samples passing all six KPIs"))
    rows.append(("SDR (generation level)", f'=COUNTIF(B4:B{last},"CLEAN")/{len(sheet_names)}',
                 "Clean generations divided by total generations"))
    rows.append(("Majority vote outcome",
                 f'=IF(COUNTIF(B4:B{last},"CLEAN")>={len(sheet_names) / 2},"CLEAN","NOT CLEAN")',
                 "Scenario-level binary used for the paired McNemar tests"))
    rows.append(("Output stable across samples?",
                 f'=IF(COUNTIF(B4:B{last},B4)={len(sheet_names)},"identical class",'
                 f'"class varies between samples")',
                 "If the class varies, intra-prompt variance is present - report it"))

    r = last + 2
    style_cell(ws.cell(r, 1, "Derived measures"), BOLD, HDR_FILL)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    for label, formula, note in rows:
        r += 1
        style_cell(ws.cell(r, 1, label), BOLD)
        ws.cell(r, 2, formula)
        style_cell(ws.cell(r, 2), BASE, align="center")
        style_cell(ws.cell(r, 5, note), wrap=True)
    ws.cell(last + 4, 2).number_format = "0.0%"

    r += 2
    style_cell(ws.cell(r, 1, "How to read this workbook"), BOLD, HDR_FILL)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    legend = [
        "One sheet per sample. Column B holds the KPI verdict, column C the failure detail.",
        "Yellow cells are the only ones to edit by hand; everything else is written by the pipeline.",
        "Class in B9 of each run sheet is a formula - do not overwrite it.",
        "K4 records N/A when the manifest creates no Pods (vacuity rule); N/A counts as a pass.",
        "K6 fails when any one scanner reports a HIGH or CRITICAL canonical fault after dedup.",
    ]
    for line in legend:
        r += 1
        style_cell(ws.cell(r, 1, line), wrap=True)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    return ws


def findings_sheet(wb, results):
    ws = wb.create_sheet("Findings")
    heads = ["Sample", "Canonical fault", "Severity", "Resource",
             "Scanners agreeing", "Rule ID", "Title"]
    widths = [12, 30, 12, 34, 20, 26, 52]
    for i, (h, w) in enumerate(zip(heads, widths), start=1):
        style_cell(ws.cell(1, i, h), BOLD, HDR_FILL)
        ws.column_dimensions[get_column_letter(i)].width = w

    r = 2
    for label, res in results:
        for f in (res or {}).get("findings", []):
            vals = [label, f.get("canonical_fault"), f.get("severity"),
                    f.get("resource"), ", ".join(f.get("scanners", [])),
                    f.get("rule_id"), f.get("title")]
            for i, v in enumerate(vals, start=1):
                style_cell(ws.cell(r, i, v), wrap=(i == 7))
            r += 1
    if r == 2:
        style_cell(ws.cell(2, 1, "No findings recorded yet - run evaluate.py first."))
    return ws


def build(template=False, out=None):
    wb = Workbook()
    wb.remove(wb.active)

    if template:
        results = [(f"Run {i}", None) for i in (1, 2, 3)]
    else:
        files = sorted((ROOT / "results").glob("*.json"))
        if not files:
            raise SystemExit("no results/*.json found - run evaluate.py first, "
                             "or pass --template")
        results = []
        for i, f in enumerate(files, start=1):
            results.append((f"Run {i}", json.loads(f.read_text())))

    names = [label for label, _ in results]
    for label, res in results:
        run_sheet(wb, label, res)
    summary_sheet(wb, names)
    findings_sheet(wb, results)
    wb.move_sheet("Summary", offset=-len(wb.sheetnames) + 1)

    out = out or (ROOT / "S01_P0_M1_pilot.xlsx")
    wb.save(out)
    print(f"wrote {out}")
    return out


if __name__ == "__main__":
    p = argparse.ArgumentParser()
    p.add_argument("--template", action="store_true")
    p.add_argument("--out")
    a = p.parse_args()
    build(a.template, a.out)
