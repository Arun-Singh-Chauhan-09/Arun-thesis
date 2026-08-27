#!/usr/bin/env python3
"""
Render ONE generation's K1..K6 outcome into the Excel layout:

            B                C
    1       Passed/Failed    Summary
    2  KPI1
    ...
    7  KPI6

Sheet is named YAML1 (YAML2, YAML3, ... if you pass more than one run).

Usage:
    python3 scripts/make_yaml_sheet.py                       # all runs found
    python3 scripts/make_yaml_sheet.py S01_P0_M1_r1          # one run
    python3 scripts/make_yaml_sheet.py -o KPI_result.xlsx
"""
import argparse
import json
import pathlib
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = pathlib.Path(__file__).resolve().parent.parent
RESULTS = ROOT / "results"

KPIS = ["K1", "K2", "K3", "K4", "K5", "K6"]
KPI_QUESTION = {
    "K1": "Parseable YAML (PyYAML)",
    "K2": "Valid Kubernetes schema (kubeconform)",
    "K3": "Cluster accepts it (kubectl apply)",
    "K4": "Becomes ready (kubectl wait, 180s)",
    "K5": "Does what was asked (intent probe)",
    "K6": "No HIGH/CRITICAL faults (Checkov/Trivy/kube-linter)",
}

HEAD = Font(name="Arial", size=11, bold=True)
BODY = Font(name="Arial", size=11)
PASS_FILL = PatternFill("solid", fgColor="C6EFCE")
FAIL_FILL = PatternFill("solid", fgColor="FFC7CE")
NA_FILL = PatternFill("solid", fgColor="EDEDED")
PASS_FONT = Font(name="Arial", size=11, bold=True, color="006100")
FAIL_FONT = Font(name="Arial", size=11, bold=True, color="9C0006")


def dig(obj, key):
    """Find `key` anywhere in a nested dict/list. Returns None if absent."""
    if isinstance(obj, dict):
        if key in obj:
            return obj[key]
        for v in obj.values():
            found = dig(v, key)
            if found is not None:
                return found
    elif isinstance(obj, list):
        for v in obj:
            found = dig(v, key)
            if found is not None:
                return found
    return None


def kpi_state(record, k):
    """Return (verdict, summary_text) for one KPI."""
    raw = dig(record, k)

    # KPI may be stored as a bare bool, or as {"passed": ..., "detail": ...}
    detail = ""
    if isinstance(raw, dict):
        detail = str(
            raw.get("detail") or raw.get("reason") or raw.get("message") or ""
        )
        raw = raw.get("passed", raw.get("result", raw.get("value")))

    if raw is None or (isinstance(raw, str) and raw.upper() == "N/A"):
        return "N/A", detail or "Not applicable for this manifest"
    verdict = "Passed" if bool(raw) else "Failed"
    return verdict, detail


def k6_summary(record):
    """Build a fault summary line for K6 from whatever findings were stored."""
    findings = dig(record, "findings") or dig(record, "faults") or []
    if not isinstance(findings, list) or not findings:
        return ""
    sev = {}
    ids = []
    for f in findings:
        if not isinstance(f, dict):
            continue
        s = str(f.get("severity", "")).upper()
        if s:
            sev[s] = sev.get(s, 0) + 1
        rid = f.get("rule_id") or f.get("id")
        if rid and str(rid) not in ids:
            ids.append(str(rid))
    parts = []
    if sev:
        parts.append(
            ", ".join(f"{n} {s}" for s, n in sorted(sev.items()))
        )
    if ids:
        shown = ", ".join(ids[:6])
        parts.append(f"({shown}{', ...' if len(ids) > 6 else ''})")
    return " ".join(parts)


def build_sheet(ws, record, run_id):
    ws["A1"] = run_id
    ws["A1"].font = HEAD
    ws["B1"] = "Passed/Failed"
    ws["B1"].font = HEAD
    ws["C1"] = "Summary"
    ws["C1"].font = HEAD

    for i, k in enumerate(KPIS, start=2):
        verdict, detail = kpi_state(record, k)

        label = ws.cell(row=i, column=1, value=f"KPI{k[1]}")
        label.font = BODY

        cell = ws.cell(row=i, column=2, value=verdict)
        cell.alignment = Alignment(horizontal="center")
        if verdict == "Passed":
            cell.font, cell.fill = PASS_FONT, PASS_FILL
        elif verdict == "Failed":
            cell.font, cell.fill = FAIL_FONT, FAIL_FILL
        else:
            cell.font, cell.fill = BODY, NA_FILL

        summary = detail
        if k == "K6" and verdict == "Failed":
            extra = k6_summary(record)
            summary = f"{extra} - {detail}".strip(" -") if extra else detail
        if not summary:
            summary = KPI_QUESTION[k]

        s = ws.cell(row=i, column=3, value=summary)
        s.font = BODY
        s.alignment = Alignment(vertical="top", wrap_text=True)

    # Final classification, two rows below the ladder
    verdict = dig(record, "classification") or dig(record, "class") or ""
    ws["A9"] = "Classification"
    ws["A9"].font = HEAD
    ws["B9"] = str(verdict)
    ws["B9"].font = HEAD
    ws["C9"] = "INVALID = fails K1/K2 | MISCONFIGURED = fails any of K3-K6 | CLEAN = passes K1-K6"
    ws["C9"].font = BODY

    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 15
    ws.column_dimensions["C"].width = 78
    ws.freeze_panes = "A2"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("runs", nargs="*", help="run ids, e.g. S01_P0_M1_r1")
    ap.add_argument("-o", "--out", default="KPI_result.xlsx")
    a = ap.parse_args()

    if not RESULTS.is_dir():
        sys.exit(f"no results directory at {RESULTS}")

    files = sorted(RESULTS.glob("*.json"))
    if a.runs:
        files = [f for f in files if any(r in f.stem for r in a.runs)]
    if not files:
        sys.exit(f"no matching result JSON in {RESULTS}")

    wb = Workbook()
    wb.remove(wb.active)

    for n, f in enumerate(files, start=1):
        record = json.loads(f.read_text())
        ws = wb.create_sheet(f"YAML{n}")
        build_sheet(ws, record, f.stem)
        found = [k for k in KPIS if dig(record, k) is not None]
        print(f"YAML{n}  {f.name}  KPIs found: {found or 'NONE - check JSON shape'}")

    out = ROOT / a.out
    wb.save(out)
    print(f"wrote {out}")


if __name__ == "__main__":
    main()