#!/usr/bin/env python3
"""
Turn one evaluate.py result into the simple KPI table in Excel.

Reads results/<gen_id>.json (running evaluate.py first if it is missing) and
writes a single sheet: KPI | What it checks | passed/failed/skipped | Details.

Usage:
    python3 scripts/kpi_report.py --gen S01_P0_M1_manual
    python3 scripts/kpi_report.py --gen S01_P0_M1_manual -o kpi_results.xlsx
    python3 scripts/kpi_report.py --gen S01_P0_M1_manual --rerun
"""
import argparse
import json
import pathlib
import subprocess
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = pathlib.Path(__file__).resolve().parent.parent
RESULTS = ROOT / "results"

KPI_QUESTIONS = [
    ("KPI1", "K1", "Parseable YAML (PyYAML)"),
    ("KPI2", "K2", "Valid Kubernetes schema (kubeconform)"),
    ("KPI3", "K3", "Cluster accepts it (kubectl apply)"),
    ("KPI4", "K4", "Becomes ready (kubectl wait, 180s)"),
    ("KPI5", "K5", "Does what was asked (intent probe)"),
    ("KPI6", "K6", "No HIGH/CRITICAL faults (Checkov/Trivy/kube-linter)"),
]

# --- STYLES ---
HEADER_FONT = Font(name="Calibri", size=11, bold=True)
BODY_FONT = Font(name="Calibri", size=11)
OK_FONT = Font(name="Calibri", size=11, color="006100")
BAD_FONT = Font(name="Calibri", size=11, color="9C0006")
OK_FILL = PatternFill("solid", start_color="C6EFCE", end_color="C6EFCE")
BAD_FILL = PatternFill("solid", start_color="FFC7CE", end_color="FFC7CE")
SKIP_FILL = PatternFill("solid", start_color="EDEDED", end_color="EDEDED")
HEADER_FILL = PatternFill("solid", start_color="DCE6F1", end_color="DCE6F1")
BORDER = Border(left=Side(style="thin"), right=Side(style="thin"),
                top=Side(style="thin"), bottom=Side(style="thin"))
CENTER = Alignment(horizontal="center", vertical="center")
WRAP = Alignment(vertical="center", wrap_text=True)


def rung(rec, k):
    """Normalise one rung to True / False / None, whatever shape it is in."""
    v = rec.get("verdict", rec)
    raw = v.get(k) if isinstance(v, dict) else None
    if isinstance(raw, dict):
        raw = raw.get("passed", raw.get("result"))
    return raw if isinstance(raw, bool) else None


def load(gen_id, rerun):
    """Return the result record, running evaluate.py if needed."""
    path = RESULTS / f"{gen_id}.json"
    if rerun or not path.exists():
        r = subprocess.run(
            [sys.executable, str(ROOT / "scripts" / "evaluate.py"), "--gen", gen_id],
            text=True, cwd=ROOT,
        )
        if not path.exists():
            sys.exit(f"evaluate.py produced no {path} (exit {r.returncode})")
    return json.loads(path.read_text())


def main():
    ap = argparse.ArgumentParser(description=__doc__.strip().split("\n")[0])
    ap.add_argument("--gen", required=True, help="generation id, e.g. S01_P0_M1_manual")
    ap.add_argument("-o", "--out", default=None, help="output .xlsx (default: <gen>_kpi.xlsx)")
    ap.add_argument("--rerun", action="store_true", help="re-run evaluate.py even if cached")
    a = ap.parse_args()

    rec = load(a.gen, a.rerun)
    detail = rec.get("detail", {}) or {}
    classification = rec.get("classification", rec.get("class", "?"))

    # where the ladder stopped, used to explain the skipped rungs
    blocked_at = next((k for _, k, _ in KPI_QUESTIONS if rung(rec, k) is False), None)

    wb = Workbook()
    ws = wb.active
    ws.title = "KPI Results"

    for col, header in enumerate(["KPI", "What it checks",
                                  "passed/failed/skipped", "Details/Errors"], start=1):
        c = ws.cell(row=1, column=col, value=header)
        c.font, c.fill, c.border, c.alignment = HEADER_FONT, HEADER_FILL, BORDER, CENTER

    for i, (label, k, question) in enumerate(KPI_QUESTIONS, start=2):
        c = ws.cell(row=i, column=1, value=label)
        c.font, c.border, c.alignment = HEADER_FONT, BORDER, CENTER

        c = ws.cell(row=i, column=2, value=question)
        c.font, c.border, c.alignment = BODY_FONT, BORDER, WRAP

        state = rung(rec, k)
        if state is True:
            text, font, fill = "passed", OK_FONT, OK_FILL
        elif state is False:
            text, font, fill = "failed", BAD_FONT, BAD_FILL
        else:
            text, font, fill = "skipped", BODY_FONT, SKIP_FILL
        c = ws.cell(row=i, column=3, value=text)
        c.font, c.fill, c.border, c.alignment = font, fill, BORDER, CENTER

        msg = detail.get(k) or ""
        if not msg:
            if state is True:
                msg = "check passed"
            elif state is None and blocked_at:
                msg = f"Not attempted - {classification} at {blocked_at}"
            elif state is None:
                msg = "Not attempted"
        c = ws.cell(row=i, column=4, value=str(msg))
        c.font, c.border, c.alignment = BODY_FONT, BORDER, WRAP

    row = len(KPI_QUESTIONS) + 3
    ws.cell(row=row, column=1, value="Generation").font = HEADER_FONT
    ws.cell(row=row, column=2, value=a.gen).font = BODY_FONT
    ws.cell(row=row + 1, column=1, value="Classification").font = HEADER_FONT
    ws.cell(row=row + 1, column=2, value=str(classification)).font = HEADER_FONT
    ws.cell(row=row + 2, column=1, value="Rule").font = HEADER_FONT
    ws.cell(row=row + 2, column=2,
            value="INVALID = fails K1/K2 | MISCONFIGURED = fails any of K3-K6 "
                  "| CLEAN = passes K1-K6").font = BODY_FONT

    for col, w in (("A", 12), ("B", 45), ("C", 22), ("D", 60)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"

    out = ROOT / (a.out or f"{a.gen}_kpi.xlsx")
    wb.save(out)

    summary = "  ".join(f"{k}={rung(rec, k)}" for _, k, _ in KPI_QUESTIONS)
    print(f"{a.gen}  {classification}  {summary}")
    print(f"wrote {out}")


if __name__ == "__main__":
    main()